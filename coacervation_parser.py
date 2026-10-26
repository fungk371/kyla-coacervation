"""
Coacervation kinetic-plate parser.

Handles the wide-format kinetic plate-reader export used for coacervation studies:

    Time    Temperature(C)  A1  A2  ... H12
    00:00:00    37          92.0 92.3 ... 89.6
    00:01:00    37          ...
    ...

Differences from the original single-surfactant transmission format:
  - Wide layout: wells are COLUMNS (A1..H12), each row is a TIMEPOINT.
  - One file holds ALL timepoints (not one file per timepoint).
  - Time is an HH:MM:SS string in a column (not a decimal in the filename).
  - A Temperature column is included.
  - Plain tab-delimited text (not UTF-16).

This module melts that grid into the long format the database expects:
  well, row, column, timepoint_minutes, temperature_c, transmission_pct
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pandas as pd


WELL_RE = re.compile(r"^([A-H])(\d{1,2})$")


def _read_raw(file_path: str | Path) -> pd.DataFrame:
    """Read the tab-delimited kinetic file, tolerating encoding + preambles.

    Handles two export dialects:
      * plain: the first line is the 'Time  Temperature  A1..H12' header.
      * SoftMax-style: a '##BLOCKS=' / 'Plate:' preamble precedes the header,
        and trailing 'Group:' blocks follow a '~End' marker (both stripped).
    Also detects UTF-16 (BOM) in addition to utf-8 / latin-1.
    """
    data = Path(file_path).read_bytes()
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        text = data.decode("utf-16")
    elif data[:3] == b"\xef\xbb\xbf":
        text = data.decode("utf-8-sig")
    else:
        for enc in ("utf-8", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover
            raise ValueError(f"Could not decode {file_path} with known encodings")

    lines = text.splitlines()

    # Find the real header row (first line beginning with 'Time').
    start = next((i for i, ln in enumerate(lines)
                  if ln.strip().lower().startswith("time")), 0)
    # Stop at the end-of-block / group summaries marker.
    end = len(lines)
    for j in range(start + 1, len(lines)):
        s = lines[j].strip()
        if s.startswith("~End") or s.startswith("Group") or s.startswith("##"):
            end = j
            break
    block = "\n".join(lines[start:end])

    df = pd.read_csv(io.StringIO(block), sep="\t")
    # Drop trailing all-empty columns produced by a trailing tab.
    df = df.dropna(axis=1, how="all")
    # Strip whitespace from column names.
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _time_to_minutes(value: str) -> float:
    """Convert an 'HH:MM:SS' string or a datetime.time to float minutes."""
    import datetime as _dt
    if isinstance(value, _dt.time):
        return value.hour * 60.0 + value.minute + value.second / 60.0
    if isinstance(value, _dt.datetime):
        return value.hour * 60.0 + value.minute + value.second / 60.0
    parts = str(value).strip().split(":")
    parts = [float(p) for p in parts]
    if len(parts) == 3:
        h, m, s = parts
    elif len(parts) == 2:
        h, m, s = 0.0, parts[0], parts[1]
    else:
        raise ValueError(f"Unrecognized time value: {value!r}")
    return h * 60.0 + m + s / 60.0


def _sheet_is_kinetic(ws) -> bool:
    """True if a worksheet's first row looks like Time / Temp / A1..H12."""
    try:
        header = [str(ws.cell(row=1, column=c).value or "").strip()
                  for c in range(1, 20)]
    except Exception:  # noqa: BLE001
        return False
    has_time = any(h.lower().startswith("time") for h in header)
    has_well = any(WELL_RE.match(h) for h in header)
    return has_time and has_well


def find_kinetic_sheet(path) -> "str | None":
    """Return the name of the sheet holding the raw kinetic grid, if any.

    Prefers a sheet whose name mentions 'plate reader' / 'plate 1' / 'plate',
    else the first sheet whose header looks like Time + wells.
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return None
    names = list(wb.sheetnames)
    ordered = sorted(names, key=lambda n: (
        0 if "plate reader" in n.lower() else
        1 if n.lower().strip() in ("plate 1", "plate1") else
        2 if "plate" in n.lower() else 3))
    hit = None
    for n in ordered:
        if _sheet_is_kinetic(wb[n]):
            hit = n
            break
    wb.close()
    return hit


def parse_plate1_sheet(file_path, sheet: str = None) -> "pd.DataFrame":
    """Parse the raw-kinetic tab of a result .xlsm (raw plate reader grid pasted
    in), auto-detecting the sheet ('Plate 1', 'plate reader', or any tab with a
    Time + A1..H12 header). Time may be a datetime.time. Returns long format.
    """
    from openpyxl import load_workbook
    if sheet is None:
        sheet = find_kinetic_sheet(file_path)
    wb = load_workbook(file_path, data_only=True, read_only=True)
    if sheet is None or sheet not in wb.sheetnames:
        sheet = next((n for n in wb.sheetnames if _sheet_is_kinetic(wb[n])), wb.sheetnames[0])
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError(f"'{sheet}' sheet is empty in {file_path}")
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    time_i = next((i for i, h in enumerate(header) if h.lower().startswith("time")), 0)
    temp_i = next((i for i, h in enumerate(header) if "temp" in h.lower()), 1)
    well_idx = [(i, h) for i, h in enumerate(header) if WELL_RE.match(h)]
    if not well_idx:
        raise ValueError("No well columns (A1..H12) found in 'Plate 1' header")

    long_rows = []
    filename = Path(file_path).name
    for r in rows[1:]:
        if time_i >= len(r) or r[time_i] is None:
            continue
        try:
            tmin = _time_to_minutes(r[time_i])
        except ValueError:
            continue
        try:
            temp_c = float(r[temp_i]) if temp_i < len(r) and r[temp_i] is not None else None
        except (TypeError, ValueError):
            temp_c = None
        for i, wc in well_idx:
            if i >= len(r) or r[i] is None:
                continue
            try:
                pct = float(r[i])
            except (TypeError, ValueError):
                continue
            m = WELL_RE.match(wc)
            long_rows.append({
                "well": wc, "row": m.group(1), "column": int(m.group(2)),
                "timepoint_minutes": round(tmin, 4), "temperature_c": temp_c,
                "transmission_pct": pct, "transmission_filename": filename,
            })
    long_df = pd.DataFrame(long_rows)
    if long_df.empty:
        raise ValueError(f"No data parsed from '{sheet}' in {file_path}")
    return long_df


def parse_plate1_sample_info(file_path, sheet: str = None) -> dict:
    """Extract the sample-info block that sits BELOW the kinetic grid on a
    result .xlsm (Trough Position | BatchID | SampleID | Description | Hardness
    | HHS Shake ...). Returns {"1": {...}, "2": {...}} keyed by trough order."""
    from openpyxl import load_workbook
    if sheet is None:
        sheet = find_kinetic_sheet(file_path)
    if sheet is None:
        return {}
    wb = load_workbook(file_path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return {}
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()

    hdr_i, colmap = None, {}
    for i, r in enumerate(rows):
        low = [str(c).strip().lower() if c is not None else "" for c in r]
        if "batchid" in low and "description" in low:
            hdr_i = i
            colmap = {low[j]: j for j in range(len(low)) if low[j]}
            break
    if hdr_i is None:
        return {}

    def get(r, key):
        j = colmap.get(key)
        return r[j] if j is not None and j < len(r) else None

    out, order = {}, 0
    for r in rows[hdr_i + 1:]:
        batch, desc = get(r, "batchid"), get(r, "description")
        if not batch and not desc:
            continue
        order += 1
        out[str(order)] = {
            "batch_id": str(batch or "").strip(),
            "sample_id": str(get(r, "sampleid") or "").strip(),
            "description": str(desc or "").strip(),
            "hardness": str(get(r, "hardness") or "").strip(),
            "shake_speed": get(r, "hhs shake speed (rpm)"),
        }
        if order >= 2:
            break
    return out



def parse_kinetic_plate_file(file_path: str | Path) -> pd.DataFrame:
    """Parse a wide-format kinetic coacervation file into long format.

    Args:
        file_path: Path to the tab-delimited kinetic export.

    Returns:
        DataFrame with columns:
            well, row, column, timepoint_minutes, temperature_c,
            transmission_pct, transmission_filename
    """
    raw = _read_raw(file_path)

    if raw.shape[1] < 3:
        raise ValueError(
            f"Expected Time, Temperature and well columns; got {list(raw.columns)}"
        )

    # Identify the time and temperature columns (first two by convention).
    time_col = raw.columns[0]
    temp_col = raw.columns[1]
    if "time" not in time_col.lower():
        raise ValueError(f"First column should be Time, found {time_col!r}")

    # Well columns are everything matching a well pattern (A1..H12).
    well_cols = [c for c in raw.columns if WELL_RE.match(c)]
    if not well_cols:
        raise ValueError("No well columns (A1..H12) found in file header")

    long_rows = []
    filename = Path(file_path).name
    for _, r in raw.iterrows():
        try:
            tmin = _time_to_minutes(r[time_col])
        except ValueError:
            continue  # skip blank/footer rows
        try:
            temp_c = float(r[temp_col])
        except (TypeError, ValueError):
            temp_c = None

        for wc in well_cols:
            val = r[wc]
            if pd.isna(val):
                continue
            m = WELL_RE.match(wc)
            long_rows.append(
                {
                    "well": wc,
                    "row": m.group(1),
                    "column": int(m.group(2)),
                    "timepoint_minutes": round(tmin, 4),
                    "temperature_c": temp_c,
                    "transmission_pct": float(val),
                    "transmission_filename": filename,
                }
            )

    long_df = pd.DataFrame(long_rows)
    if long_df.empty:
        raise ValueError(f"No data parsed from {file_path}")
    return long_df


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python coacervation_parser.py <kinetic_file.txt>")
        raise SystemExit(1)

    parsed = parse_kinetic_plate_file(sys.argv[1])
    print(parsed.head(20).to_string(index=False))
    print()
    print(f"Rows: {len(parsed)}")
    print(f"Wells: {parsed['well'].nunique()}")
    print(f"Timepoints (min): {sorted(parsed['timepoint_minutes'].unique())}")
    print(
        f"Temperature range: "
        f"{parsed['temperature_c'].min()}-{parsed['temperature_c'].max()} C"
    )
    print(
        f"%T range: {parsed['transmission_pct'].min():.2f} - "
        f"{parsed['transmission_pct'].max():.2f}"
    )
