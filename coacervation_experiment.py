"""
Build a Nexus submission workbook (flex-driver format) for a coacervation
experiment - e.g. a 3 chassis x 8 polymer shampoo design.

Mirrors the standardized template in the repo-root make_spreadsheets.py
(nexus_flex_driver_template.xlsx) but FILLS it for every sample in the
experiment, so the whole set is submitted as one worksheet.

Tabs produced:
  1_Experiment       - experiment-level metadata (worksheet id, submitter, ...)
  2_Samples_Prep     - one row per sample: Nexus key (BatchID) + all prep params
  3_Raw_Data_XY      - long X-Y raw curve per sample -> nexus_xy
  4_Nexus_Critical   - long flex-driver rows (sample_id, test_method, param, value)

Also emits a flat nexus_xy EAV CSV matching cdl_stc_prod.silver_nexus.nexus_xy.

The workbook builder is pure (no Flask / parsing) - callers pass a list of
`sample` dicts already computed from the plate data.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------- styling
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
NEXUS_FILL = PatternFill("solid", fgColor="E2EFDA")   # green = goes to Nexus
XY_FILL = PatternFill("solid", fgColor="DDEBF7")      # blue = goes to XY table
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(italic=True, color="808080", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TEST_METHOD_DEFAULT = "Coacervation Dilution (R&D)"
PARAM_LIST_ID = "Coacervation Data"


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    if note:
        ws["A2"] = note
        ws["A2"].font = NOTE_FONT


def _autosize(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --------------------------------------------------------------------------- #
# Prep parameters that travel with every sample (all "critical" per current
# guidance).  (parameter_name, sample_key, unit)
# --------------------------------------------------------------------------- #
PREP_PARAMS = [
    ("chassis", "chassis", "-"),
    ("polymer", "polymer", "-"),
    ("polymer_concentration", "concentration", "% w/w"),
    ("surfactant", "surfactant", "-"),
    ("polymer_mw", "mw_kda", "kDa"),
    ("polymer_charge_density", "charge", "meq/g"),
    ("water_hardness", "hardness", "gpg"),
    ("shake_speed", "shake_speed_rpm", "rpm"),
    ("shake_temperature", "shake_temp_c", "degC"),
    ("shake_time", "shake_time_s", "sec"),
    ("plate_temperature", "temperature_c", "degC"),
]

# Coacervation critical results.  (parameter_name, sample_key, unit)
COAC_PARAMS = [
    ("coac_auc_absorbance", "auc", "A*fraction"),
    ("coac_onset_dilution", "onset", "ratio_1:X"),
    ("coac_peak_dilution", "peak", "ratio_1:X"),
    ("coac_trough_dilution", "trough_dilution", "ratio_1:X"),
    ("coac_transition_width", "width", "-"),
    ("coac_recovery", "recovery", "-"),
    ("coac_strength", "strength", "-"),
    ("coac_timepoint", "timepoint", "min"),
]


def critical_rows_for(sample: dict, test_method: str) -> list[dict]:
    """Flatten a sample into flex-driver critical rows (prep + coacervation)."""
    out = []
    # experiment role (variant / benchmark / control) + optional control name
    for name, key in (("sample_role", "role"), ("control_name", "control_name")):
        val = sample.get(key)
        if val:
            out.append({
                "sample_id": sample.get("sample_id", ""),
                "test_method": test_method,
                "parameter_name": name,
                "result_value": val,
                "unit": "-",
                "flag/notes": "",
            })
    for name, key, unit in PREP_PARAMS + COAC_PARAMS:
        val = sample.get(key)
        if val is None or val == "":
            continue
        out.append({
            "sample_id": sample.get("sample_id", ""),
            "test_method": test_method,
            "parameter_name": name,
            "result_value": val,
            "unit": unit,
            "flag/notes": sample.get("flag", ""),
        })
    return out


def build_nexus_workbook(
    exp_meta: dict,
    samples: list[dict],
    out_dir: str | Path = "exports",
    *,
    test_method: Optional[str] = None,
) -> dict:
    """Write the filled flex-driver workbook + flat nexus_xy CSV.

    Args:
        exp_meta: experiment-level metadata (name, worksheet_id, submitter, ...).
        samples:  list of sample dicts (see module docstring / route builder).
        out_dir:  directory to write outputs into.
        test_method: overrides the default R&D method label.

    Returns dict with xlsx_path, csv_path, critical_rows (preview), n_samples.
    """
    test_method = test_method or exp_meta.get("test_method") or TEST_METHOD_DEFAULT
    out = Path(out_dir)
    out.mkdir(exist_ok=True)
    safe = "".join(ch if (ch.isalnum() or ch in "-_") else "_"
                   for ch in str(exp_meta.get("name") or "experiment")).strip("_") or "experiment"

    wb = Workbook()

    # ---- Tab 1: Experiment metadata ----
    ws = wb.active
    ws.title = "1_Experiment"
    _title(ws, "Coacervation Experiment - Nexus submission",
           "Applies to all samples listed on the following tabs.")
    ws.append([]); ws.append(["Field", "Value"])
    _style_header(ws, 3, 2)
    meta_rows = [
        ("Experiment name", exp_meta.get("name", "")),
        ("Worksheet / Request ID", exp_meta.get("worksheet_id", "")),
        ("Submitter", exp_meta.get("submitter", "")),
        ("Submission date", exp_meta.get("date", "")),
        ("Study / Project", exp_meta.get("study", "")),
        ("Test method", test_method),
        ("Instrument", exp_meta.get("instrument", "Plate reader (%T)")),
        ("Purpose / Request description", exp_meta.get("purpose", "")),
        ("GxP status", exp_meta.get("gxp", "Non-GxP")),
        ("Replicates", exp_meta.get("replicates", "One and done")),
        ("Timepoint (min)", exp_meta.get("timepoint", "")),
        ("AUC dilution window", exp_meta.get("window", "")),
        ("Design", exp_meta.get("design", "")),
        ("N samples", len(samples)),
        ("Raw data location", "cdl_stc_prod.silver_nexus.nexus_xy"),
    ]
    for f, v in meta_rows:
        ws.append([f, v])
        for c in (1, 2):
            ws.cell(row=ws.max_row, column=c).border = BORDER
    _autosize(ws, [30, 55])

    # ---- Tab 2: Samples & prep parameters (one row per sample) ----
    ws = wb.create_sheet("2_Samples_Prep")
    _title(ws, "Samples & Preparation Parameters (all critical)",
           "One row per sample. sample_id = Nexus sample key (BatchID).")
    ws.append([])
    hdr = ["sample_id", "role", "control_name", "run_tab", "chassis", "polymer",
           "concentration", "surfactant", "mw_kda", "charge_meq_g", "hardness",
           "shake_speed_rpm", "shake_temp_C", "shake_time_s", "plate_temp_C"]
    ws.append(hdr)
    _style_header(ws, 3, len(hdr))
    for c in range(1, len(hdr) + 1):
        ws.cell(row=3, column=c).fill = NEXUS_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="375623")
    for s in samples:
        ws.append([
            s.get("sample_id", ""), s.get("role", ""), s.get("control_name", ""),
            s.get("run", ""), s.get("chassis", ""),
            s.get("polymer", ""), s.get("concentration", ""), s.get("surfactant", ""),
            s.get("mw_kda"), s.get("charge"), s.get("hardness"),
            s.get("shake_speed_rpm"), s.get("shake_temp_c"), s.get("shake_time_s"),
            s.get("temperature_c"),
        ])
        for c in range(1, len(hdr) + 1):
            ws.cell(row=ws.max_row, column=c).border = BORDER
    _autosize(ws, [20, 10, 18, 14, 16, 12, 12, 12, 10, 12, 10, 14, 12, 12, 12])

    # ---- Tab 3: Raw data (long X-Y) ----
    ws = wb.create_sheet("3_Raw_Data_XY")
    _title(ws, "Raw Data  ->  cdl_stc_prod.silver_nexus.nexus_xy",
           "Long X-Y: dilution ratio vs %T and Absorbance, at the chosen timepoint.")
    ws.append([])
    hdr = ["sample_id", "series", "x_label", "x_value", "y_label", "y_value", "unit"]
    ws.append(hdr)
    _style_header(ws, 3, len(hdr))
    for c in range(1, len(hdr) + 1):
        ws.cell(row=3, column=c).fill = XY_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="1F4E78")
    for s in samples:
        sid = s.get("sample_id", "")
        for pt in s.get("curve", []):
            ws.append([sid, "dilution_curve", "dilution_ratio", pt.get("dilution"),
                       "transmittance", pt.get("pct"), "%T"])
            ws.append([sid, "dilution_curve", "dilution_ratio", pt.get("dilution"),
                       "absorbance", pt.get("absorbance"), "A"])
    _autosize(ws, [20, 16, 14, 12, 14, 12, 8])

    # ---- Tab 4: Nexus critical (flex driver) ----
    ws = wb.create_sheet("4_Nexus_Critical")
    _title(ws, "Nexus Critical Parameters  ->  Flexible Driver import",
           "Long format. One block of parameters per sample.")
    ws.append([])
    hdr = ["sample_id", "test_method", "parameter_name", "result_value", "unit", "flag/notes"]
    ws.append(hdr)
    _style_header(ws, 3, len(hdr))
    for c in range(1, len(hdr) + 1):
        ws.cell(row=3, column=c).fill = NEXUS_FILL
        ws.cell(row=3, column=c).font = Font(bold=True, color="375623")
    critical_rows = []
    for s in samples:
        for cr in critical_rows_for(s, test_method):
            critical_rows.append(cr)
            ws.append([cr["sample_id"], cr["test_method"], cr["parameter_name"],
                       cr["result_value"], cr["unit"], cr["flag/notes"]])
            for c in range(1, len(hdr) + 1):
                ws.cell(row=ws.max_row, column=c).border = BORDER
    _autosize(ws, [20, 26, 24, 14, 12, 18])

    xlsx_path = out / f"nexus_flex_{safe}.xlsx"
    wb.save(xlsx_path)

    # ---- flat nexus_xy EAV CSV ----
    csv_path = out / f"nexus_xy_{safe}.csv"
    eav = []
    for s in samples:
        for cr in critical_rows_for(s, test_method):
            eav.append({
                "key_id1": cr["sample_id"],
                "param_list_id": PARAM_LIST_ID,
                "param_list_version_id": 1,
                "variant_id": f"{s.get('chassis', '')}|{s.get('polymer', '')}",
                "dataset": 1,
                "replicate_id": 1,
                "param_id": cr["parameter_name"],
                "entered_value": cr["result_value"],
                "entered_units": cr["unit"],
                "_source_file": s.get("file", ""),
            })
    pd.DataFrame(eav).to_csv(csv_path, index=False)

    return {
        "xlsx_path": str(xlsx_path),
        "csv_path": str(csv_path),
        "critical_rows": critical_rows,
        "n_samples": len(samples),
        "test_method": test_method,
    }
