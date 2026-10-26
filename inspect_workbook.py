"""Inspect the coacervation formulation workbooks to locate the formula fraction."""

from pathlib import Path
import openpyxl

RAW = Path("raw_data")
files = [
    "2025 1024 COMPARE GH130 GH150.xlsm",
    "2025 1023 01.xlsm",
]

for fname in files:
    path = RAW / fname
    if not path.exists():
        print(f"MISSING: {path}")
        continue
    print("=" * 70)
    print(f"FILE: {fname}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    print("Sheets:", wb.sheetnames)

    # Look for the '96 well plate calculations' sheet (fuzzy match)
    target = None
    for s in wb.sheetnames:
        if "96" in s and "calc" in s.lower():
            target = s
            break
    if target is None:
        # fallback: any sheet containing 'plate'
        for s in wb.sheetnames:
            if "plate" in s.lower() or "well" in s.lower():
                target = s
                break
    print("Target sheet:", target)

    if target:
        ws = wb[target]
        # Dump the header area around column AI and the AI3:AI50 range
        print("\n-- Column AI, rows 1..52 (formula fraction area) --")
        for r in range(1, 53):
            val = ws[f"AI{r}"].value
            # also grab a couple of neighbor columns for context
            ah = ws[f"AH{r}"].value
            aj = ws[f"AJ{r}"].value
            print(f"row {r:>2}: AH={ah!r:>20}  AI={val!r:>20}  AJ={aj!r!r:>20}")
    wb.close()
