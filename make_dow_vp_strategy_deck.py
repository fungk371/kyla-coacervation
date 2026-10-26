from pathlib import Path
from shutil import copyfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "Reports"
OUT = REPORTS / "DOW_Dextran_Band5_VP_Strategy_Deck.pptx"
OUT_COPY = Path(r"C:\VScode_Projects\DOW_Dextran_Band5_VP_Strategy_Deck.pptx")
FONT = "Open Sans"

BG = RGBColor(0xFA, 0xFC, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x13, 0x2E, 0x4C)
LIGHT_NAVY = RGBColor(0x36, 0x4D, 0x6B)
DUSTY = RGBColor(0x66, 0x82, 0xA5)
PASTEL = RGBColor(0xD8, 0xE7, 0xF2)
LINE = RGBColor(0xC9, 0xD8, 0xE6)
TEAL = RGBColor(0x7E, 0xA3, 0xBF)
GREEN = RGBColor(0xA4, 0xC0, 0xD6)
GOLD = RGBColor(0xBE, 0xD2, 0xE4)
CORAL = RGBColor(0x4E, 0x6F, 0x94)
GREY = RGBColor(0x5D, 0x6B, 0x7A)
DARK = RGBColor(0x1D, 0x35, 0x52)
PALE_PANEL = RGBColor(0xE8, 0xF0, 0xF6)
MIST = RGBColor(0xEF, 0xF6, 0xFB)

ASSETS = {
    "chem": ROOT / "_dextran_docx_extract" / "unzipped" / "word" / "media" / "image2.png",
    "auc": REPORTS / "vp_dow_auc_window.png",
    "window": REPORTS / "dextran_use_windows.png",
    "conc": REPORTS / "dextran_conc_response.png",
    "salt": REPORTS / "salt_vs_nosalt.png",
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def set_font(run, size=14, bold=False, italic=False, color=NAVY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def tx(slide, left, top, width, height, items, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    frame.margin_top = Pt(1)
    frame.margin_bottom = Pt(1)
    for idx, item in enumerate(items):
        text, opts = item if isinstance(item, tuple) else (item, {})
        para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        para.alignment = align
        para.space_after = Pt(opts.get("space_after", 5))
        run = para.add_run()
        run.text = ("•  " if opts.get("bullet") else "") + text
        set_font(run, opts.get("size", 14), opts.get("bold", False), opts.get("italic", False), opts.get("color", NAVY))
    return box


def fill_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    soft_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.20), SW, Inches(1.30))
    soft_band.fill.solid()
    soft_band.fill.fore_color.rgb = MIST
    soft_band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), SH)
    accent.fill.solid()
    accent.fill.fore_color.rgb = DUSTY
    accent.line.fill.background()


def title(slide, text, kicker):
    head = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.22))
    head.fill.solid()
    head.fill.fore_color.rgb = BG
    head.line.fill.background()
    tx(slide, Inches(0.55), Inches(0.12), Inches(8.2), Inches(0.38), [(kicker.upper(), {"size": 14, "bold": True, "color": DUSTY})])
    tx(slide, Inches(0.55), Inches(0.48), Inches(12.0), Inches(0.70), [(text, {"size": 26, "bold": True, "color": LIGHT_NAVY})])
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.10), Inches(0.68), Pt(4))
    rule.fill.solid()
    rule.fill.fore_color.rgb = DUSTY
    rule.line.fill.background()


def footer(slide, number, takeaway):
    label = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(6.34), Inches(2.18), Inches(0.78))
    label.fill.solid()
    label.fill.fore_color.rgb = DUSTY
    label.line.fill.background()
    tx(slide, Inches(0.74), Inches(6.53), Inches(1.78), Inches(0.28), [("BUSINESS IMPACT", {"size": 14, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    foot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.48), Inches(6.34), Inches(10.30), Inches(0.78))
    foot.fill.solid()
    foot.fill.fore_color.rgb = LIGHT_NAVY
    foot.line.fill.background()
    tx(slide, Inches(2.76), Inches(6.48), Inches(8.96), Inches(0.42), [(takeaway, {"size": 15, "bold": True, "color": WHITE})], anchor=MSO_ANCHOR.MIDDLE)
    tx(slide, Inches(11.88), Inches(6.55), Inches(0.62), Inches(0.28), [(f"{number}/10", {"size": 10, "color": PASTEL})], align=PP_ALIGN.RIGHT)


def new_slide(text, number, takeaway, kicker="COACERVATION STRATEGY"):
    slide = prs.slides.add_slide(BLANK)
    fill_bg(slide)
    title(slide, text, kicker)
    footer(slide, number, takeaway)
    return slide


def card(slide, left, top, width, height, heading, body, accent=DUSTY, fill=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.25)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(5))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    items = []
    if heading:
        items.append((heading, {"size": 15.5, "bold": True, "color": LIGHT_NAVY, "space_after": 8}))
    items.extend(body)
    tx(slide, left + Inches(0.18), top + Inches(0.20), width - Inches(0.36), height - Inches(0.30), items)


def impact_tag(slide, left, top, text_value, color=DUSTY):
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.25), Inches(0.42))
    tag.fill.solid()
    tag.fill.fore_color.rgb = color
    tag.line.fill.background()
    tx(slide, left + Inches(0.10), top + Inches(0.08), Inches(2.05), Inches(0.22), [(text_value, {"size": 14, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)


def pill(slide, left, top, width, text, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.38))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tx(slide, left + Inches(0.06), top + Inches(0.04), width - Inches(0.12), Inches(0.28), [(text, {"size": 14, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)


def fit(slide, path, left, top, width, height):
    path = Path(path)
    if not path.exists():
        card(slide, left, top, width, height, "Image pending", [(path.name, {"size": 14, "color": CORAL})], CORAL)
        return
    with Image.open(path) as img:
        image_width, image_height = img.size
    aspect = image_width / image_height
    box_aspect = width / height
    pic_width, pic_height = (width, int(width / aspect)) if aspect >= box_aspect else (int(height * aspect), height)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE
    frame.line.width = Pt(1)
    slide.shapes.add_picture(path.as_posix(), left + int((width - pic_width) / 2), top + int((height - pic_height) / 2), width=pic_width, height=pic_height)


def hamilton_slot(slide, left, top, width, height):
    card(slide, left, top, width, height, "Hamilton STAR automation", [], DUSTY)
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.45), top + Inches(0.86), width - Inches(0.90), Inches(1.82))
    body.fill.solid()
    body.fill.fore_color.rgb = DARK
    body.line.fill.background()
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.72), top + Inches(1.10), width - Inches(1.44), Inches(0.09))
    rail.fill.solid()
    rail.fill.fore_color.rgb = TEAL
    rail.line.fill.background()
    head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(3.0), top + Inches(1.30), Inches(0.82), Inches(0.58))
    head.fill.solid()
    head.fill_fore_color = PASTEL
    head.fill.fore_color.rgb = PASTEL
    head.line.fill.background()
    for col in range(12):
        for row in range(8):
            well = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.88 + col * 0.15), top + Inches(2.03 + row * 0.062), Inches(0.052), Inches(0.052))
            well.fill.solid()
            well.fill.fore_color.rgb = PASTEL if (row + col) % 3 else GOLD
            well.line.fill.background()
    tx(slide, left + Inches(0.40), top + Inches(3.0), width - Inches(0.80), Inches(0.72), [("48 dilution conditions per sample • time-resolved %T • standardized raw files", {"size": 14, "bold": True})], align=PP_ALIGN.CENTER)
    tx(slide, left + Inches(0.40), top + Inches(3.64), width - Inches(0.80), Inches(0.44), [("Photo slot: replace with your pasted Hamilton image once saved as a file.", {"size": 14, "italic": True, "color": GREY})], align=PP_ALIGN.CENTER)


def platform_mockup(slide, left, top, width, height):
    card(slide, left, top, width, height, "Coacervation analysis website", [], DUSTY)
    labels = [("Analyze %T", DUSTY), ("Phase maps", TEAL), ("Lasentec", GREEN), ("Composition", LIGHT_NAVY)]
    positions = [(0.34, 0.82), (2.18, 0.82), (0.34, 1.28), (2.18, 1.28)]
    for (label, color), (xoff, yoff) in zip(labels, positions):
        pill(slide, left + Inches(xoff), top + Inches(yoff), Inches(1.62), label, color)
    plot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.38), top + Inches(1.88), width - Inches(0.76), Inches(1.10))
    plot.fill.solid()
    plot.fill.fore_color.rgb = RGBColor(0xF8, 0xFB, 0xFD)
    plot.line.color.rgb = LINE
    for idx, y_offset in enumerate([0.25, 0.35, 0.62, 0.48, 0.22, 0.16, 0.33, 0.60, 0.80]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.62 + idx * 0.34), top + Inches(2.00 + y_offset), Inches(0.06), Inches(0.06))
        dot.fill.solid()
        dot.fill.fore_color.rgb = CORAL
        dot.line.fill.background()
    card(slide, left + Inches(0.38), top + Inches(3.25), width - Inches(0.76), Inches(0.95), "Decision metrics", [("AUC • onset • consumer-window score • convergence", {"size": 14, "bold": True})], GREEN)


def workflow_step(slide, left, top, width, number, label, color):
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.42), Inches(0.42))
    circ.fill.solid()
    circ.fill.fore_color.rgb = color
    circ.line.fill.background()
    tx(slide, left, top + Inches(0.08), Inches(0.42), Inches(0.20), [(number, {"size": 14, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)
    tx(slide, left + Inches(0.52), top + Inches(0.02), width - Inches(0.52), Inches(0.38), [(label, {"size": 14, "bold": True, "color": NAVY})], anchor=MSO_ANCHOR.MIDDLE)


def feature_chip(slide, left, top, width, heading, body, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.72))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = color
    shape.line.width = Pt(1.25)
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.12), top + Inches(0.16), Inches(0.34), Inches(0.34))
    icon.fill.solid()
    icon.fill.fore_color.rgb = color
    icon.line.fill.background()
    tx(slide, left + Inches(0.58), top + Inches(0.10), width - Inches(0.70), Inches(0.22), [(heading, {"size": 14, "bold": True, "color": LIGHT_NAVY})])
    tx(slide, left + Inches(0.58), top + Inches(0.36), width - Inches(0.70), Inches(0.22), [(body, {"size": 14, "color": GREY})])


def phase_map(slide, left, top, width, height):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE
    tx(slide, left + Inches(0.18), top + Inches(0.16), width - Inches(0.36), Inches(0.28), [("Phase diagram map", {"size": 15, "bold": True, "color": LIGHT_NAVY})])
    chart_left, chart_top = left + Inches(0.42), top + Inches(0.78)
    chart_w, chart_h = width - Inches(0.84), height - Inches(1.28)
    area = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, chart_left, chart_top, chart_w, chart_h)
    area.fill.solid()
    area.fill.fore_color.rgb = RGBColor(0xF8, 0xFB, 0xFD)
    area.line.color.rgb = LINE
    for i, color in enumerate([PASTEL, RGBColor(0xC7, 0xD9, 0xE5), TEAL, DUSTY]):
        band = slide.shapes.add_shape(MSO_SHAPE.ARC, chart_left + Inches(0.30 + i * 0.18), chart_top + Inches(0.18 + i * 0.10), chart_w - Inches(0.50 + i * 0.30), chart_h - Inches(0.10 + i * 0.30))
        band.line.color.rgb = color
        band.line.width = Pt(5)
    pill(slide, chart_left + Inches(1.15), chart_top + Inches(1.0), Inches(2.35), "coacervate window", DUSTY)
    tx(slide, chart_left + Inches(0.12), chart_top + chart_h - Inches(0.36), Inches(1.8), Inches(0.28), [("polymer", {"size": 14, "color": GREY})])
    tx(slide, chart_left + chart_w - Inches(2.25), chart_top + Inches(0.06), Inches(2.05), Inches(0.28), [("surfactant / hardness", {"size": 14, "color": GREY})], align=PP_ALIGN.RIGHT)


def slide_1():
    s = prs.slides.add_slide(BLANK)
    fill_bg(s)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.24), SH)
    bar.fill.solid()
    bar.fill.fore_color.rgb = DUSTY
    bar.line.fill.background()
    tx(s, Inches(0.72), Inches(0.48), Inches(8.2), Inches(0.42), [("BAND 5 VP STRATEGY PRESENTATION", {"size": 14, "bold": True, "color": DUSTY})])
    tx(s, Inches(0.72), Inches(1.04), Inches(11.7), Inches(1.55), [("Accelerating Superior Shampoo Innovation", {"size": 36, "bold": True, "color": LIGHT_NAVY})])
    tx(s, Inches(0.75), Inches(2.38), Inches(10.4), Inches(0.54), [("Building the Digital Formulation-Learning Engine for Hair Care", {"size": 23, "bold": True, "color": DUSTY})])
    tx(s, Inches(0.75), Inches(3.18), Inches(10.55), Inches(1.02), [("Accelerating product innovation through high-throughput experimentation, digital decision-making, and new approaches to formulation science that enable faster learning, faster innovation cycles, and superior products for consumers.", {"size": 16})])
    card(s, Inches(0.75), Inches(4.22), Inches(3.85), Inches(1.38), "Strategic shift", [("From: Does this coacervate?", {"size": 14.5, "color": GREY}), ("To: When, where, and why does it create the right mechanism?", {"size": 14.5, "bold": True})], DUSTY)
    card(s, Inches(4.90), Inches(4.22), Inches(3.55), Inches(1.38), "Business goal", [("Accelerate sustainable formulation decisions before late-stage rework.", {"size": 14.5, "bold": True})], TEAL)
    card(s, Inches(8.75), Inches(4.22), Inches(3.55), Inches(1.38), "First proof point", [("DOW dextran screening as the first real application.", {"size": 14.5, "bold": True})], GREEN)
    footer(s, 1, "Automation becomes strategic when it accelerates the path from analytical data to absolute-superiority choices.")


slide_1()

s = new_slide("Agenda: From Business Need to Future Formulation Strategy", 2, "The deck links business need, platform capability, DOW evidence, phase maps, and future strategy.", "EXECUTIVE ROADMAP")
agenda = [
    ("1", "Business need", "Faster learning cycles for sustainable polymers, new surfactants, and superior products."),
    ("2", "Capability built", "Digital workflow converting raw %T and phase-diagram files into reusable decision metrics."),
    ("3", "First application", "DOW dextran as a real sustainable-polymer screening and formulation-learning case."),
    ("4", "Performance linkage", "Phase diagrams, Lasentec, microscopy, rheology, deposition, and Sumanth wet-conditioning/T2P work."),
    ("5", "Future vision", "A predictive coacervation learning engine for faster absolute superiority."),
]
for y, (num, head, body) in zip([1.42, 2.28, 3.14, 4.00, 4.86], agenda):
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.82), Inches(y), Inches(0.50), Inches(0.50))
    circ.fill.solid(); circ.fill.fore_color.rgb = DUSTY; circ.line.fill.background()
    tx(s, Inches(0.82), Inches(y + 0.08), Inches(0.50), Inches(0.22), [(num, {"size": 14, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)
    tx(s, Inches(1.55), Inches(y - 0.02), Inches(10.6), Inches(0.58), [(head, {"size": 16.5, "bold": True, "color": LIGHT_NAVY}), (body, {"size": 14, "color": GREY})])

s = new_slide("The Business Need: Faster Learning Cycles for Superior Products", 3, "The business risk is learning too late which polymer-formulation combinations can win.", "BUSINESS NEED")
card(s, Inches(0.65), Inches(1.40), Inches(3.85), Inches(4.05), "Need 1: Faster innovation", [("Sustainable polymers and new surfactants are not drop-in replacements.", {"size": 14.5}), ("Accelerate by screening more formulation conditions earlier.", {"size": 14.5, "bold": True, "color": LIGHT_NAVY})], DUSTY)
card(s, Inches(4.75), Inches(1.40), Inches(3.85), Inches(4.05), "Need 2: Digital decisions", [("Raw plate-reader files and complex curves must become stored metrics teams can compare.", {"size": 14.5}), ("Accelerate by turning data into reusable decision assets.", {"size": 14.5, "bold": True, "color": LIGHT_NAVY})], TEAL)
card(s, Inches(8.85), Inches(1.40), Inches(3.85), Inches(4.05), "Need 3: Performance linkage", [("Absolute superiority depends on connecting coacervation behavior to wet conditioning and consumer feel.", {"size": 14.5}), ("Accelerate by integrating %T, phase diagrams, Lasentec, microscopy, rheology, and T2P.", {"size": 14.5, "bold": True, "color": LIGHT_NAVY})], GREEN)

s = new_slide("The Strategic Gap: Link Coacervate Properties to Wet Conditioning", 4, "Coacervate strategy must connect amount, phase behavior, size, morphology, rheology, deposition, and consumer feel.", "COACERVATE PROPERTIES")
props = [
    ("Amount", "%T / AUC", "How much optical coacervation forms across dilution.", DUSTY),
    ("Phase Behavior", "Phase diagrams", "Where formulation space creates or suppresses coacervates.", TEAL),
    ("Size", "Lasentec / FBRM", "Particle-size fingerprints during dilution and rinse.", GREEN),
    ("Morphology", "Microscopy", "What structures form and how they carry actives.", GOLD),
    ("Rheology Gap", "future method", "Missing bridge from structure to what consumers feel.", CORAL),
]
for idx, (head, sub, body, color) in enumerate(props):
    card(s, Inches(0.55 + idx * 2.45), Inches(1.45), Inches(2.18), Inches(3.15), head, [(sub, {"size": 14, "bold": True, "color": LIGHT_NAVY}), (body, {"size": 14})], color)
card(s, Inches(0.80), Inches(4.92), Inches(11.65), Inches(0.95), "Sumanth wet-conditioning/T2P linkage", [("The next strategic proof point is determining which analytical fingerprints predict deposition, conditioning, and consumer-relevant feel.", {"size": 15, "bold": True})], LIGHT_NAVY)

s = new_slide("The Platform: Turning Coacervation Data Into Formulation Decisions", 5, "The website makes coacervation data reusable across %T, phase diagrams, Lasentec, composition, and future rheology.", "DIGITAL DATA ANALYSIS PLATFORM")
platform_mockup(s, Inches(0.65), Inches(1.35), Inches(4.25), Inches(4.55))
card(s, Inches(5.20), Inches(1.35), Inches(7.30), Inches(1.10), "Screen recording cue", [("Show the website moving from raw upload to plots, metrics, saved comparisons, and cross-method context.", {"size": 15, "bold": True})], DUSTY, MIST)
workflow_step(s, Inches(5.35), Inches(2.68), Inches(2.25), "1", "raw plate files", DUSTY)
workflow_step(s, Inches(7.55), Inches(2.68), Inches(2.25), "2", "%T fingerprints", TEAL)
workflow_step(s, Inches(9.95), Inches(2.68), Inches(2.25), "3", "decision metrics", GREEN)
workflow_step(s, Inches(5.35), Inches(3.26), Inches(2.25), "4", "database/library", LIGHT_NAVY)
workflow_step(s, Inches(7.55), Inches(3.26), Inches(2.25), "5", "Nexus export", DUSTY)
workflow_step(s, Inches(9.95), Inches(3.26), Inches(2.25), "6", "PD guidance", TEAL)
feature_chip(s, Inches(5.25), Inches(4.12), Inches(3.45), "Analyze", "plots + AUC + onset", DUSTY)
feature_chip(s, Inches(8.95), Inches(4.12), Inches(3.45), "Phase", "precipitated wells + t95", TEAL)
feature_chip(s, Inches(5.25), Inches(4.98), Inches(3.45), "Lasentec", "size / floc fingerprints", GREEN)
feature_chip(s, Inches(8.95), Inches(4.98), Inches(3.45), "Composition", "Mintel / INCI context", LIGHT_NAVY)

s = new_slide("Approach: Automated Fingerprints and Phase Maps Accelerate the Next Experiment", 6, "The experiment tested polymer architecture, chassis, concentration, salt, and dilution timing as decision levers.", "APPROACH + HYPOTHESES")
for idx, (label, color) in enumerate([("Manual method", DUSTY), ("Hamilton %T", TEAL), ("Python metrics", LIGHT_NAVY), ("Phase maps", GREEN), ("PD guidance", GOLD)]):
    pill(s, Inches(0.78 + idx * 2.34), Inches(1.36), Inches(2.15), label, color)
hamilton_slot(s, Inches(0.72), Inches(2.25), Inches(5.65), Inches(3.68))
card(s, Inches(6.72), Inches(2.25), Inches(5.78), Inches(3.68), "Hypotheses tested", [("H1: One dextran architecture will separate from weaker candidates and justify focused investment.", {"size": 14.3, "bullet": True}), ("H2: Chassis chemistry will change both amount and timing of coacervation.", {"size": 14.3, "bullet": True}), ("H3: More polymer will not universally improve response; compatibility will dominate.", {"size": 14.3, "bullet": True}), ("H4: Phase maps identify formulation regions most likely to translate to wet-conditioning hypotheses.", {"size": 14.3, "bullet": True})], TEAL)

s = new_slide("First Application: DOW Dextran Screening Accelerated Lead Down-Selection", 7, "Dex-K emerged as the strongest developmental dextran, focusing the next learning cycle.", "DOW DEXTRAN RESULTS")
fit(s, ASSETS["auc"], Inches(0.65), Inches(1.35), Inches(7.70), Inches(4.78))
card(s, Inches(8.65), Inches(1.35), Inches(3.85), Inches(4.78), "What the data says", [("CD2 and N-Hance 3196 set benchmark behavior.", {"size": 14.3, "bullet": True}), ("Dex-K was the lead developmental dextran across the screen.", {"size": 14.3, "bold": True, "bullet": True}), ("Dex-A showed measurable but weaker response.", {"size": 14.3, "bullet": True}), ("B/J variants were generally near baseline at 0.4% polymer.", {"size": 14.3, "bullet": True}), ("At 1%, Dex-K approached ~80% of CD2 benchmark response in LAPB chassis.", {"size": 14.3, "bold": True, "bullet": True})], TEAL)

s = new_slide("Finding: Chassis, Salt, and Timing Determine Whether a Polymer Can Win", 8, "Performance is a polymer-formulation interaction that must activate at the right point in consumer dilution.", "FORMULATION CONTEXT")
fit(s, ASSETS["window"], Inches(0.65), Inches(1.35), Inches(5.95), Inches(4.65))
fit(s, ASSETS["salt"], Inches(6.88), Inches(1.35), Inches(5.60), Inches(4.65))
pill(s, Inches(0.90), Inches(5.62), Inches(5.40), "APG shifted Dex-K toward a later, potentially consumer-relevant window", GREEN)
pill(s, Inches(7.15), Inches(5.62), Inches(5.05), "NaCl changed Dex-K/APG optical response ~2.5x", DUSTY)

s = new_slide("Beyond DOW: Phase Diagrams Build the Bridge to Wet-Conditioning Strategy", 9, "Phase diagrams turn coacervation from a yes/no screen into a formulation map for Sumanth wet-conditioning/T2P hypotheses.", "PHASE DIAGRAM + T2P PLAN")
phase_map(s, Inches(0.65), Inches(1.35), Inches(5.50), Inches(4.70))
card(s, Inches(6.45), Inches(1.35), Inches(6.05), Inches(4.70), "Experimental plan logic", [("%T fingerprints: whether, when, and how strongly coacervation forms during dilution.", {"size": 14.2, "bullet": True}), ("Phase diagrams: where polymer, surfactant, hardness, and chassis conditions create the coacervate region.", {"size": 14.2, "bullet": True}), ("Lasentec/microscopy/rheology: what structure forms, how large it is, and how it may feel.", {"size": 14.2, "bullet": True}), ("Wet conditioning/T2P: test whether analytical fingerprints predict consumer-relevant performance.", {"size": 14.2, "bold": True, "bullet": True})], TEAL)

s = new_slide("Future Vision: A Predictive Coacervation Learning Engine for Absolute Superiority", 10, "Beauty Analytical can identify what works, understand why it works, and determine what to test next, faster.", "FUTURE STRATEGY")
card(s, Inches(0.65), Inches(1.35), Inches(3.80), Inches(2.0), "Accelerate down-selection", [("Prioritize winning polymers and surfactants earlier using reusable fingerprints.", {"size": 14.5, "bold": True})], DUSTY)
card(s, Inches(4.75), Inches(1.35), Inches(3.80), Inches(2.0), "Accelerate formulation design", [("Use chassis, salt, hardness, concentration, and phase maps as designed levers.", {"size": 14.5, "bold": True})], TEAL)
card(s, Inches(8.85), Inches(1.35), Inches(3.80), Inches(2.0), "Accelerate prediction", [("Link %T, phase diagrams, Lasentec, microscopy, rheology, deposition, and wet conditioning.", {"size": 14.5, "bold": True})], GREEN)
card(s, Inches(0.65), Inches(3.70), Inches(12.0), Inches(2.0), "Strategic roadmap", [("Near term: standardize %T and phase-diagram metrics in the website/database.", {"size": 14.5, "bullet": True}), ("Mid term: integrate cross-method characterization and composition context.", {"size": 14.5, "bullet": True}), ("Long term: predictive database for faster polymer-surfactant-chassis selection and absolute superiority.", {"size": 14.5, "bold": True, "bullet": True})], LIGHT_NAVY)

prs.save(OUT)
try:
    copyfile(OUT, OUT_COPY)
except PermissionError:
    print(f"Could not copy to {OUT_COPY}; file may be open. Saved primary copy at {OUT}")
print(f"Saved {OUT}")
if OUT_COPY.exists():
    print(f"Copied {OUT_COPY}")
print(f"Slides: {len(prs.slides)}")