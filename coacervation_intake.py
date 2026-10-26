"""
Coacervation intake workbook: a clean, metadata-only template (one row per
sample) that pairs with the raw .txt plate exports.

The user fills a single workbook describing every sample in a study. Each row
links to a raw plate file via `Plate ID` (the raw .txt filename) + `Trough`
(1 or 2). Raw %T still comes from the wide instrument .txt export.

MW / charge density entered here OVERRIDE component_properties.csv.

Public API:
  build_intake_template(path)     -> write the blank template .xlsx
  parse_intake_workbook(path)     -> list[dict] (one per sample row)
  build_intake_store(rows)        -> {norm_plate: {"1": {...}, "2": {...}}}
  norm_plate(s)                   -> normalized join key
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------- #
# Column definitions: (key, label, required, help, options)
# `key` is the internal field name; `label` is the header shown in the sheet.
# `options` (if set) becomes a dropdown data-validation.
# --------------------------------------------------------------------------- #
COLUMNS = [
    ("plate_id", "Plate ID (raw .txt file)", True,
     "Must match the raw plate file name you upload (extension optional).", None),
    ("trough", "Trough (1 or 2)", True,
     "1 = wells A-D, 2 = wells E-H.", ["1", "2"]),
    ("batch_id", "Batch ID", True,
     "Unique physical-sample key; used as the Nexus sample key (key_id1).", None),
    ("sample_id", "Sample ID", False, "Lab sample identifier.", None),
    ("name", "Sample Name / Description", False,
     "Human-readable name, e.g. 'GH150 EED0 + 0.4% Dex-K'.", None),
    ("chassis", "Chassis", False, "Base formulation, e.g. 'GH150 EED0'.", None),
    ("variable", "Changing (polymer/surfactant)", False,
     "Which component is being varied in this study.", ["polymer", "surfactant"]),
    ("polymer", "Polymer", False, "Polymer name/grade, e.g. '3196', 'Dex-K'.", None),
    ("poly_mw", "Polymer MW (kDa)", False, "Measured polymer molecular weight.", None),
    ("poly_cd", "Polymer charge density (meq/g)", False,
     "Measured polymer charge density.", None),
    ("surfactant", "Surfactant", False, "Surfactant name, e.g. 'SLES'.", None),
    ("surf_mw", "Surfactant MW (kDa)", False, "Surfactant molecular weight (if relevant).", None),
    ("surf_cd", "Surfactant charge density (meq/g)", False,
     "Surfactant charge density (if relevant).", None),
    ("concentration", "Concentration", False, "e.g. '0.4%'.", None),
    ("hardness", "Water hardness", False, "e.g. '7 gpg'.", None),
    ("particles_gt_30um", "Coacervate particles > 30 um", False,
     "Lasentec/FBRM count of coacervate flocs above 30 um (drives PO depo).", None),
    ("coacervate_size_um", "Coacervate size (um)", False,
     "Characteristic coacervate floc size, if measured.", None),
    ("shake_speed", "Shake speed (rpm)", False, "HHS shake speed.", None),
    ("shake_temp", "Shake temp (C)", False, "HHS shake temperature.", None),
    ("shake_time", "Shake time (s)", False, "HHS shake time.", None),
    ("replicate", "Replicate", False, "Replicate id/number (optional).", None),
    ("ph", "pH", False, "Sample pH (optional).", None),
    ("salt_mM", "Salt / ionic strength (mM)", False, "Optional.", None),
    ("prep_date", "Prep date", False, "Optional.", None),
    ("submitter", "Submitter", False, "Optional.", None),
    ("notes", "Notes", False, "Free text (optional).", None),
]

HEADER_ROW = 1
DATA_START_ROW = 2

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REQ_FILL = PatternFill("solid", fgColor="FCE4D6")   # light orange = required
TITLE_FONT = Font(bold=True, size=13, color="1F4E78")
NOTE_FONT = Font(italic=True, color="808080", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm_plate(s) -> str:
    """Normalize a Plate ID / filename to a stable join key."""
    s = str(s or "").strip()
    s = re.sub(r"\.(txt|xlsx|xlsm|csv)$", "", s, flags=re.IGNORECASE)
    s = re.sub(r"\s*\(\d+\)\s*$", "", s)   # strip download duplicate suffix ' (1)'
    return s.strip().lower()


def plate_code(name) -> str:
    """Extract the trough-pair code used to pair a raw .txt with its metadata
    sheet, e.g. 'SDA_2025.1106_07_08.txt' and 'in 20251106 07 08 ...xlsx' both
    -> '07_08'. Controls -> 'controls'."""
    s = str(name).lower()
    s = re.sub(r"\.(txt|xlsx|xlsm|csv)$", "", s)
    s = re.sub(r"\(\d+\)", "", s)
    s = re.sub(r"20\d{2}[.]?\d{4}", " ", s)   # strip date 20251106 / 2025.1106
    matches = re.findall(r"(\d{2})[ _]+(\d{2})", s)
    if matches:
        a, b = matches[-1]
        return f"{a}_{b}"
    if "control" in s:
        return "controls"
    toks = [t for t in re.split(r"[ _]+", s) if t]
    return toks[-1] if toks else s


def parse_description(desc, component: str = "") -> dict:
    """Split a formula Description into chassis / varying component / %.

    'GH150 16% R2510 + 10% Tea Saponin' -> chassis 'GH150 16% R2510',
    component 'Tea Saponin', concentration '10%'.
    If `component` is given, its % is extracted (tolerating abbreviations);
    otherwise the first '<n>% <name>' after the +/- separator is auto-detected.
    """
    d = str(desc or "").strip()
    if not d:
        return {"chassis": "", "component": "", "concentration": ""}
    chassis = re.split(r"\s*\+\s*|\s*\u2013\s*|\s+-\s+", d, maxsplit=1)[0].strip()
    comp = str(component or "").strip()
    conc, comp_name = "", comp
    if comp:
        m = re.search(r"([\d.]+)\s*%\s*" + re.escape(comp), d, re.I)
        if not m and comp.split():
            m = re.search(r"([\d.]+)\s*%\s*" + re.escape(comp.split()[0]), d, re.I)
        if m:
            conc = m.group(1) + "%"
    else:
        m = re.search(r"[+\u2013\-]\s*([\d.]+)\s*%\s*([A-Za-z][\w /]*)", d)
        if m:
            conc = m.group(1) + "%"
            comp_name = m.group(2).strip()
    return {"chassis": chassis, "component": comp_name, "concentration": conc}


def parse_metadata_sheet(path, component: str = "", comp_type: str = "surfactant") -> list[dict]:
    """Parse a flat metadata sheet (Trough Position, BatchID, SampleID,
    Description, Hardness, HHS Shake ...) into up to 2 sample records, with the
    Description auto-split into chassis / component / concentration."""
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr_row = None
    for r in range(1, 6):
        vals = [str(ws.cell(row=r, column=c).value or "").strip().lower()
                for c in range(1, ws.max_column + 1)]
        if "batchid" in vals or "description" in vals:
            hdr_row = r
            break
    if hdr_row is None:
        wb.close()
        raise ValueError("No BatchID / Description header found in metadata sheet.")
    colmap = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(row=hdr_row, column=c).value or "").strip().lower()
        if h:
            colmap[h] = c

    def get(row, key):
        c = colmap.get(key)
        return ws.cell(row=row, column=c).value if c else None

    samples, order, r = [], 0, hdr_row + 1
    while r <= ws.max_row and order < 2:
        batch, desc = get(r, "batchid"), get(r, "description")
        if not batch and not desc:
            r += 1
            continue
        order += 1
        p = parse_description(desc, component)
        rec = {
            "trough": "1" if order == 1 else "2",
            "batch_id": str(batch or "").strip(),
            "sample_id": str(get(r, "sampleid") or "").strip(),
            "name": str(desc or "").strip(),
            "description": str(desc or "").strip(),
            "chassis": p["chassis"], "concentration": p["concentration"],
            "variable": comp_type,
            "hardness": str(get(r, "hardness") or "").strip(),
            "shake_speed": get(r, "hhs shake speed (rpm)"),
            "shake_time": get(r, "hhs shake time (sec)"),
            "shake_temp": get(r, "hhs shake temp (c)"),
        }
        if comp_type == "surfactant":
            rec["surfactant"] = p["component"]
        else:
            rec["polymer"] = p["component"]
        samples.append(rec)
        r += 1
    wb.close()
    return samples



# --------------------------------------------------------------------------- #
# Template generation
# --------------------------------------------------------------------------- #
def build_intake_template(path: str | Path) -> str:
    """Write a blank intake workbook (Samples sheet + Instructions sheet)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"

    # header row
    for c, (key, label, required, _help, _opts) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(26, len(label) + 2))

    # two example rows (a benchmark + a variant on the same plate)
    example = {
        "plate_id": "GH150_Dextrans_0805", "trough": "1", "batch_id": "B-20260722-01427",
        "sample_id": "S-0001", "name": "GH150 EED0 + 0.4% 3196", "chassis": "GH150 EED0",
        "variable": "polymer", "polymer": "3196", "poly_mw": 250, "poly_cd": 1.2,
        "surfactant": "SLES", "surf_mw": "", "surf_cd": "", "concentration": "0.4%",
        "hardness": "7 gpg", "shake_speed": 500, "shake_temp": 40, "shake_time": 60,
        "replicate": 1, "ph": "", "salt_mM": "", "prep_date": "", "submitter": "K Fung",
        "notes": "example row — delete before submitting",
    }
    example2 = dict(example, trough="2", batch_id="B-20260722-01428",
                    name="GH150 EED0 + 0.4% Dex-K", polymer="Dex-K", poly_mw=300,
                    poly_cd=1.6, replicate=1)
    for r, ex in enumerate((example, example2), start=DATA_START_ROW):
        for c, (key, _label, required, _help, _opts) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=ex.get(key, ""))
            cell.border = BORDER
            cell.font = NOTE_FONT

    # dropdowns for columns with options (apply to a generous row range)
    last = 500
    for c, (key, _label, _req, _help, opts) in enumerate(COLUMNS, start=1):
        if opts:
            col = get_column_letter(c)
            dv = DataValidation(type="list", formula1='"%s"' % ",".join(opts), allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}{DATA_START_ROW}:{col}{last}")

    ws.freeze_panes = "A2"

    # instructions sheet
    ins = wb.create_sheet("Instructions")
    ins["A1"] = "Coacervation intake template — how to use"
    ins["A1"].font = TITLE_FONT
    lines = [
        "",
        "Fill ONE row per sample on the 'Samples' sheet (delete the two example rows).",
        "",
        "Join to raw data:",
        "  • Plate ID  = the raw .txt plate file name you upload (extension optional).",
        "  • Trough    = 1 (wells A-D) or 2 (wells E-H).",
        "  Every plate has two samples, so two rows share one Plate ID (trough 1 and 2).",
        "",
        "Required columns are shaded orange: Plate ID, Trough, Batch ID.",
        "  • Batch ID is the physical-sample key and becomes the Nexus key_id1.",
        "",
        "Polymer / Surfactant MW (kDa) and charge density (meq/g):",
        "  • Enter measured values here to OVERRIDE component_properties.csv.",
        "  • 'Changing' picks which component (polymer or surfactant) is the one varied.",
        "",
        "Raw %T data still comes from the instrument .txt export (wide plate format).",
        "Upload this workbook in the app under Manage data → Intake workbook.",
        "",
        "Questions: Kyla Fung — fung.k.1@pg.com",
    ]
    for i, text in enumerate(lines, start=2):
        ins.cell(row=i, column=1, value=text).font = NOTE_FONT if i > 2 else Font(size=11)
    ins.column_dimensions["A"].width = 90

    # shade required headers
    for c, (key, _label, required, _help, _opts) in enumerate(COLUMNS, start=1):
        if required:
            ws.cell(row=HEADER_ROW, column=c).fill = PatternFill("solid", fgColor="C55A11")

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
_ALIASES = {}
for _key, _label, *_ in COLUMNS:
    _ALIASES[_label.strip().lower()] = _key
# a few tolerant aliases
_ALIASES.update({
    "plate id": "plate_id", "plate": "plate_id", "raw file": "plate_id",
    "trough": "trough", "trough position": "trough",
    "batch id": "batch_id", "batchid": "batch_id",
    "sample id": "sample_id", "sampleid": "sample_id",
    "sample name": "name", "description": "name", "name": "name",
    "changing": "variable", "variable": "variable",
    "polymer mw (kda)": "poly_mw", "polymer mw": "poly_mw",
    "polymer charge density (meq/g)": "poly_cd", "polymer charge density": "poly_cd",
    "surfactant mw (kda)": "surf_mw", "surfactant charge density (meq/g)": "surf_cd",
    "water hardness": "hardness", "hardness": "hardness",
})


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def parse_intake_workbook(path: str | Path) -> list[dict]:
    """Read the intake workbook -> list of per-sample dicts (internal keys)."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Samples"] if "Samples" in wb.sheetnames else wb[wb.sheetnames[0]]

    # find the header row (first row containing a recognizable header)
    header_row = None
    for r in range(1, 6):
        vals = [str(ws.cell(row=r, column=c).value or "").strip().lower()
                for c in range(1, ws.max_column + 1)]
        if any(v in _ALIASES for v in vals):
            header_row = r
            break
    if header_row is None:
        wb.close()
        raise ValueError("Could not find a header row (need 'Plate ID', 'Batch ID', ...).")

    # map column index -> internal key
    colmap = {}
    for c in range(1, ws.max_column + 1):
        label = str(ws.cell(row=header_row, column=c).value or "").strip().lower()
        if label in _ALIASES:
            colmap[c] = _ALIASES[label]

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        rec = {}
        for c, key in colmap.items():
            rec[key] = _clean(ws.cell(row=r, column=c).value)
        # a real sample row needs both a plate id and a batch id
        if not rec.get("plate_id") or not rec.get("batch_id"):
            continue
        rec.setdefault("variable", "polymer")
        rows.append(rec)
    wb.close()
    return rows


def build_intake_store(rows: list[dict]) -> dict:
    """Group parsed rows into {norm_plate: {"1": {...}, "2": {...}}}."""
    store = {}
    for rec in rows:
        key = norm_plate(rec.get("plate_id"))
        tr = str(rec.get("trough") or "1").strip()
        tr = "2" if tr.endswith("2") else "1"
        store.setdefault(key, {})[tr] = rec
    return store
