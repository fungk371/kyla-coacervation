"""Build the Inverse %T method technical report (.docx).

Covers: the method (Eric Johnson / SWIC inverse %T), how it was adapted to the
automated Hamilton + SpectraMax workflow, what information it delivers, and the
business impact. Embeds the method figures from Reports/.

Output: Reports/Inverse_PctT_Method_Report.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

R = "Reports"
NAVY = RGBColor(0x1F, 0x4E, 0x78)
GREY = RGBColor(0x5A, 0x64, 0x70)


def h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NAVY
    p.space_before = Pt(10); p.space_after = Pt(4)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAVY
    p.space_before = Pt(8); p.space_after = Pt(2)
    return p


def body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(10.5)
    p.space_after = Pt(6)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text); r.font.size = Pt(10.5)
    return p


def figure(doc, fname, caption):
    path = os.path.join(R, fname)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph()
        cr = c.add_run("Figure. " + caption); cr.italic = True
        cr.font.size = Pt(9); cr.font.color.rgb = GREY
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.space_after = Pt(10)


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ---- masthead ----
    t = doc.add_paragraph()
    tr = t.add_run("Technical Report")
    tr.bold = True; tr.font.size = Pt(20); tr.font.color.rgb = NAVY
    s = doc.add_paragraph()
    sr = s.add_run("An Automated Inverse-%T Coacervation Method for Opaque Shampoos\n"
                   "(Hamilton liquid handler + SpectraMax plate reader)")
    sr.font.size = Pt(13); sr.bold = True
    m = doc.add_paragraph()
    mr = m.add_run("Author: Kyla Fung, P&G Beauty Analytical  |  Method basis: Eric Johnson (SWIC) "
                   "inverse-%T coacervation profile  |  P&G Confidential")
    mr.font.size = Pt(9); mr.font.color.rgb = GREY
    doc.add_paragraph()

    # ---- 1. Executive summary / business impact ----
    h1(doc, "1. Executive summary — business impact")
    body(doc, "Coacervation is the phase separation that occurs when a shampoo is diluted during "
              "rinsing; it is the primary vehicle that deposits conditioning actives, silicone and "
              "fragrance onto hair. Being able to measure it quickly, quantitatively and at scale is a "
              "direct lever on product performance and on the speed of formulation development.")
    body(doc, "Standard %T (turbidity) coacervation profiling only works for CLEAR shampoos. A large and "
              "growing share of the portfolio — pearlised, silicone- and ZPT-containing products — is "
              "OPAQUE, where standard %T saturates and cannot see coacervation. This report describes how "
              "the inverse-%T method (which does work on opaque products) was adapted to our automated "
              "Hamilton + SpectraMax workflow, turning a manual, one-off measurement into a high-throughput, "
              "digitised capability.")
    bullet(doc, "Unlocks coacervation measurement for OPAQUE formulas — previously a blind spot for the "
                "standard method.")
    bullet(doc, "Throughput: ~96 dilution points per plate, two formulas per plate, read as a kinetic "
                "series (0–15 min) with no manual pipetting — a full polymer × chassis screen in a day.")
    bullet(doc, "Quantitative, comparable outputs (amount + width of coacervation) that rank polymers and "
                "chassis and benchmark P&G products against competitors on one axis.")
    bullet(doc, "Feeds directly into deposition (ILS), floc-size (Lasentec) and Nexus data flows, "
                "supporting formulation decisions and the biodegradable cationic-dextran program with Dow.")

    # ---- 2. Background ----
    h1(doc, "2. Background — why opaque shampoos need a different read")
    body(doc, "When a clear shampoo is diluted, coacervate forms and makes the sample turbid, so %T drops "
              "into a dip; the depth and position of that dip report how much coacervate forms and at what "
              "rinse dilution. In an OPAQUE shampoo the product is already turbid from pearliser / silicone "
              "/ ZPT particles, so a coacervation dip is invisible against the high background — standard "
              "%T saturates and gives no information.")

    # ---- 3. The method ----
    h1(doc, "3. The method — inverse %T (Eric Johnson / SWIC)")
    body(doc, "The inverse-%T method turns the opacity into the signal. The mechanism: as the opaque "
              "shampoo is diluted, the coacervate that forms TRAPS the opacifier particles and, on "
              "CENTRIFUGATION, pulls them out of suspension. The upper phase (supernatant) therefore "
              "CLEARS — %T rises into a hump. A clearer supernatant means MORE coacervate captured and "
              "removed the opacifier. Reading the centrifuged supernatant is what makes coacervation "
              "visible in an opaque product.")
    body(doc, "The method reports two quantities off the first hump (see figure):")
    bullet(doc, "AMOUNT of phase separation = the HEIGHT of the hump (how much %T rises above the opaque "
                "floor = how much opacifier the coacervate captured).")
    bullet(doc, "WIDTH of the coacervate-forming region = the span of dilution (onset → return) over which "
                "the first hump occurs = how broad the rinse window is.")
    figure(doc, "brandon_ipct_method_annotated.png",
           "Reading inverse %T. Green = width of the coacervate-forming region (first hump); "
           "yellow = amount (hump height). Left: a polymer whose coacervate re-dissolves at high "
           "dilution; right: one whose coacervate persists across the rinse.")
    body(doc, "A third, optional descriptor is whether the hump RETURNS toward the floor at high dilution "
              "(the coacervate re-dissolves and releases its payload) or PERSISTS (stays associated across "
              "the rinse). This is exploratory and beyond the two core readouts, but is informative for "
              "deposition timing.")

    # ---- 4. Automation adaptation ----
    h1(doc, "4. Adapting it to the automated Hamilton + SpectraMax workflow")
    body(doc, "The original method was a manual, low-throughput measurement. The adaptation industrialises "
              "each step:")
    bullet(doc, "Dilution series: the Hamilton liquid handler builds a 48-point Water:Shampoo dilution "
                "ladder (1:1 to ~1:539) for each of two troughs on a 96-well plate — reproducible, "
                "hands-free, and far finer than a manual series.")
    bullet(doc, "Read: the SpectraMax reads %T across the plate as a KINETIC series (0–15 min), so we "
                "capture how the signal develops over time, not just a single snapshot.")
    bullet(doc, "Centrifugation preserves the Eric-Johnson supernatant read so the hump reflects coacervate "
                "capture, not just base opacity thinning.")
    bullet(doc, "Digitisation: raw plate exports (including UTF-16 SpectraMax 'matrix' exports) are parsed "
                "automatically, mapped well→dilution, converted %T↔absorbance, smoothed, and fingerprinted "
                "(floor, onset, amount/height, width, persistence) with no manual spreadsheet work.")
    bullet(doc, "Software workflow: a web app takes the analyst from raw plate + sample input sheet → "
                "auto-matched Batch IDs and names → organise/tag → ranked coacervation outputs → export to "
                "Nexus, keeping the whole opaque (inverse %T) library separate from the clear (%T) one.")

    # ---- 5. What information it gives ----
    h1(doc, "5. What information the method delivers")
    body(doc, "For every opaque formula the method now yields a comparable coacervation fingerprint. Viewed "
              "two ways — zoomed to the coacervation region (1:1–1:100) and across the full dilution range:")
    figure(doc, "brandon_ipct_method_overlay_zoom.png",
           "Inverse-%T humps in the coacervation window (1:1–1:100) for two chassis. Higher/earlier "
           "clearing = more coacervate formed earlier in the rinse.")
    figure(doc, "brandon_ipct_method_overlay_full.png",
           "Full dilution range (1:1–1:539). At extreme dilution all samples clear by trivial dilution, "
           "so the coacervation-relevant information sits in the first hump.")
    body(doc, "The two method readouts — amount (height) and width — separate formulas that make a LOT of "
              "coacervate from those that make it over a BROAD rinse window; both matter for deposition.")
    figure(doc, "brandon_ipct_method_widthheight.png",
           "The two method outputs by polymer and chassis: amount (hump height) and width of the "
           "first coacervate-forming region.")
    body(doc, "Because the read is kinetic, we also see how each coacervate develops over the 0–15 min "
              "window — a stability/ripening axis a single timepoint hides.")
    figure(doc, "brandon_ipct_timecourse.png",
           "Time-course by sample (0–15 min). Most fingerprints are time-stable; where they drift, "
           "coacervate ripening/coarsening is occurring.")

    # ---- 6. Representative findings ----
    h1(doc, "6. Representative findings")
    bullet(doc, "Chassis gates coacervation: the same polymer gives very different amount and width in "
                "different surfactant chassis — the chassis is as important as the polymer.")
    bullet(doc, "The method cleanly ranks polymers by amount and width, and distinguishes coacervate that "
                "re-dissolves in the rinse from coacervate that persists.")
    bullet(doc, "The same pipeline benchmarks P&G in-market products against competitors on a single "
                "coacervation axis, and supports the biodegradable cationic-dextran program with Dow.")

    # ---- 7. Limitations & next steps ----
    h1(doc, "7. Limitations & next steps")
    bullet(doc, "Trivial-dilution baseline: at extreme dilution the supernatant clears regardless of "
                "coacervate. A no-polymer base + opacifier control is needed to set the true baseline and "
                "confirm the first hump is coacervate-specific.")
    bullet(doc, "Payload confirmation: ICP/GC on supernatant vs pellet across dilutions would prove the "
                "coacervate carries the opacifier (and silicone / perfume oil) down — the direct link to "
                "deposition.")
    bullet(doc, "Orthogonal validation: pair inverse-%T amount/width with Lasentec (floc size) and ILS "
                "(deposition) to confirm the ranking predicts on-hair performance.")
    bullet(doc, "Rheology: a shake vs no-shake variant of the read may give a first, high-throughput proxy "
                "for coacervate stiffness — a current capability gap — to be validated against FBRM and "
                "true rheology anchors.")

    out = os.path.join(R, "Inverse_PctT_Method_Report.docx")
    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
