"""Comprehensive, minimalist analysis deck for the 1% dextran screen.

All three chassis (base EED0, LAPB, APG) covered for:
  - 0-15 min evolution, both curve versions (through-points & best-fit)
  - absorbance vs concentration (AUC panels)
  - 0.4% vs 1% comparison
  - AUC graphs (ranking, master heatmap, AUC-over-time, concentration response)
Plus a 'why LAPB wins' slide and blank observation / microscopy / next-steps slides.

Output: Reports/Dextran_1pct_Analysis_Deck.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

R = "Reports"
FONT = "Abadi"
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x22, 0x2A, 0x33)
GREY = RGBColor(0x6B, 0x74, 0x80)
BOXBG = RGBColor(0xF4, 0xF6, 0xF9)
BOXLINE = RGBColor(0xD5, 0xDD, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]; SW, SH = prs.slide_width, prs.slide_height
_N = [1]


def _font(run, size=14, bold=False, color=INK, italic=False):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.italic = italic; run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", FONT)


def text(slide, l, t, w, h, lines, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(4)
        s, o = (ln, {}) if isinstance(ln, str) else ln
        if o.get("bullet"):
            s = "\u2022  " + s
        r = p.add_run(); r.text = s
        _font(r, size=o.get("size", 14), bold=o.get("bold", False),
              color=o.get("color", INK), italic=o.get("italic", False))
    return tb


def head(slide, title, sub=None):
    text(slide, Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.68),
         [(title, {"size": 24, "bold": True, "color": NAVY})])
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(0.98), Inches(2.4), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb = GOLD; rule.line.fill.background()
    if sub:
        text(slide, Inches(0.58), Inches(1.03), Inches(12.0), Inches(0.4),
             [(sub, {"size": 12, "color": GREY, "italic": True})])


def image(slide, path, l, t, mw, mh):
    if not os.path.exists(path):
        text(slide, l, t, mw, Inches(0.4), [(f"[missing {os.path.basename(path)}]", {"size": 10, "color": GOLD})])
        return
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih; bar = mw / mh
    if ar >= bar:
        w = mw; h = int(mw / ar)
    else:
        h = mh; w = int(mh * ar)
    slide.shapes.add_picture(path, l + (mw - w) // 2, t + (mh - h) // 2, width=w, height=h)


def obs_box(slide, l=Inches(0.55), t=Inches(6.4), w=Inches(12.25), h=Inches(0.72)):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = BOXBG; box.line.color.rgb = BOXLINE
    box.line.width = Pt(1)
    text(slide, l + Inches(0.2), t + Inches(0.08), w - Inches(0.4), h - Inches(0.16),
         [("Observations / questions:  (add here)", {"size": 11, "bold": True, "color": GREY})])


def foot(slide):
    text(slide, Inches(10.9), Inches(7.08), Inches(2.2), Inches(0.33),
         [(f"Kyla Fung  \u00b7  {_N[0]}", {"size": 8.5, "color": GREY})], align=PP_ALIGN.RIGHT)
    _N[0] += 1


def slide():
    return prs.slides.add_slide(BLANK)


def data_slide(title, sub, img, box=(0.42, 1.3, 12.5, 5.0), obs=True):
    s = slide(); head(s, title, sub)
    l, t, w, h = box
    image(s, f"{R}/{img}", Inches(l), Inches(t), Inches(w), Inches(h))
    if obs:
        obs_box(s)
    foot(s); return s


def writein(title, label):
    s = slide(); head(s, title)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.35),
                             Inches(12.25), Inches(5.55))
    box.fill.solid(); box.fill.fore_color.rgb = BOXBG; box.line.color.rgb = BOXLINE; box.line.width = Pt(1)
    text(s, Inches(0.8), Inches(1.55), Inches(11.7), Inches(0.5),
         [(label, {"size": 13, "bold": True, "color": GREY})])
    foot(s); return s


# ---- 1 title
s = slide()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.55), SW, Pt(3))
bar.fill.solid(); bar.fill.fore_color.rgb = GOLD; bar.line.fill.background()
text(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2.0), [
    ("Biodegradable Dextran \u2014 Coacervation Analysis", {"size": 36, "bold": True, "color": NAVY}),
    ("1% polymer \u00b7 3 chassis (base EED0 / LAPB / APG) \u00b7 %T vs the CD2 benchmark", {"size": 17, "color": GREY}),
])
text(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(0.5),
     [("Kyla Fung  \u00b7  Sustainable Polymers  \u00b7  September 2026", {"size": 13, "color": GREY})])
foot(s)

# ---- 2 at a glance
data_slide("At a glance \u2014 coacervation strength ranked",
           "Lead biodegradable dextrans reach ~80% of the CD2 benchmark (in LAPB)",
           "dextran_auc_ranking_all.png", box=(1.9, 1.3, 9.5, 5.0))

# ---- chassis effect at t=0 (key polymers)
data_slide("Chassis effect at t = 0 min \u2014 key polymers",
           "Freshly-mixed 3-chassis overlay (EED0 / LAPB / APG): guar, CD2, Dex-K, Dex-J, Dex-J*",
           "dextran_t0_chassis_overlay.png", box=(0.7, 1.3, 12.0, 5.0))

# ---- evolution, both curve versions, all 3 chassis
for ch, label in [("EED0", "Base (EED0)"), ("LAPB", "LAPB"), ("APG", "APG")]:
    data_slide(f"0\u201315 min evolution \u2014 {label} \u2014 through every data point",
               "Raw fidelity: honors every well (wigglier); best for seeing the late-dilution second trough",
               f"dextran_1pct_evolution_{ch}_points.png")
    data_slide(f"0\u201315 min evolution \u2014 {label} \u2014 smooth best-fit",
               "Trough-preserving smooth: clean for presenting; troughs still resolve at high dilution",
               f"dextran_1pct_evolution_{ch}_fit.png")

# ---- absorbance vs concentration (AUC panels), all 3 chassis
for ch, label in [("EED0", "Base (EED0)"), ("LAPB", "LAPB"), ("APG", "APG")]:
    data_slide(f"Absorbance AUC panels \u2014 {label}",
               "Absorbance = 2 \u2212 log10 %T (proportional to coacervate); shaded = consumer window; peak height \u2248 amount, area \u2248 load",
               f"dextran_1pct_absorbance_{ch}.png")

# ---- 0.4% vs 1%, all 3 chassis
for ch, label in [("EED0", "Base (EED0)"), ("LAPB", "LAPB"), ("APG", "APG")]:
    data_slide(f"0.4% vs 1% polymer \u2014 {label}",
               "Solid = 1%, dashed = 0.4%. More polymer deepens the trough; Dex-J* needs LAPB AND 1% to switch on",
               f"dextran_conc_compare_{ch}_fit.png")

# ---- 0.4% polymer standalone evolution (includes guar), all 3 chassis
for ch, label in [("EED0", "Base (EED0)"), ("LAPB", "LAPB"), ("APG", "APG")]:
    data_slide(f"0.4% polymer \u2014 {label} \u2014 0\u201315 min evolution",
               "Standalone 0.4% dataset (includes guar / 3196); compare shape & depth to the 1% evolution slides",
               f"dextran_04pct_evolution_{ch}.png")

# ---- AUC graphs
data_slide("All variables at a glance \u2014 AUC map",
           "Polymer \u00d7 (chassis \u00d7 concentration). Guar (3196) 1% awaiting",
           "dextran_master_heatmap.png", box=(1.7, 1.25, 10.0, 5.05))
data_slide("Coacervation AUC over time \u2014 kinetic stability",
           "Every leader loses AUC over 0\u201315 min (coarsening). Decay slope = a kinetic-stability attribute",
           "dextran_auc_over_time.png", box=(2.0, 1.3, 9.3, 5.0))
data_slide("Concentration response \u2014 0.4% vs 1%, by chassis",
           "Concentration mostly amplifies each chassis's behavior; Dex-J* (LAPB) is the exception",
           "dextran_conc_response.png", box=(0.5, 1.4, 12.3, 4.6))

# ---- why LAPB wins
s = slide(); head(s, "Why LAPB is the strongest coacervation chassis")
text(s, Inches(0.6), Inches(1.45), Inches(12.1), Inches(4.9), [
    ("Coacervation = electrostatic association of the cationic polymer with anionic surfactant micelles; strength scales with the anionic 'binding surface' the micelles present.",
     {"size": 14, "bullet": True}),
    ("LAPB (amphoteric betaine) co-micellizes with the anionic surfactant and screens headgroup repulsion \u2192 larger / wormlike micelles that bind polymer more avidly \u2192 more coacervation.",
     {"size": 14, "bullet": True, "color": NAVY, "bold": True}),
    ("The betaine can also bridge polymer and surfactant, broadening the region and shifting it to the concentrated end ('in-bottle').",
     {"size": 14, "bullet": True}),
    ("It even switches on inert polymers \u2014 Dex-J* is dead in EED0/APG but strong in LAPB.",
     {"size": 14, "bullet": True, "color": NAVY}),
    ("APG (nonionic) dilutes the anionic charge \u2192 weakest & delayed. EED0 (sulfate) is strong but less extensive than LAPB.",
     {"size": 14, "bullet": True}),
    ("Caveat: %T shows the outcome, not the mechanism \u2014 DLS / rheology / microscopy would confirm larger micelles & droplets in LAPB.",
     {"size": 12.5, "bullet": True, "italic": True, "color": GREY}),
])
foot(s)

# ---- write-in / microscopy / next steps
writein("My observations & questions", "Write your observations, interpretations, and open questions here:")

s = slide(); head(s, "Microscopy images \u2014 coacervate structure & size",
                  "Paste images here; note droplet size, liquid-vs-solid, and which formulation")
frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.45), Inches(11.9), Inches(5.35))
frame.fill.solid(); frame.fill.fore_color.rgb = WHITE; frame.line.color.rgb = BOXLINE; frame.line.width = Pt(1.25)
text(s, Inches(0.7), Inches(3.9), Inches(11.9), Inches(0.6),
     [("[ paste microscopy images here ]", {"size": 15, "italic": True, "color": BOXLINE})], align=PP_ALIGN.CENTER)
foot(s)

writein("Questions & next steps", "Next experiments, decisions needed, asks, and follow-ups:")

out = os.path.join(R, "Dextran_1pct_Analysis_Deck.pptx")
prs.save(out)
print("wrote", out, "slides:", len(prs.slides._sldIdLst))
