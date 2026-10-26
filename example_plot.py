"""
Generate an example coacervation plot from the real 'Data to Graph' sheet.

Reads formula fractions (col B) and a sample's %T column, converts to
absorbance, fits the best curve, integrates AUC over a window, and saves a PNG.
"""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from coacervation_plotting import plot_absorbance_vs_concentration

WB = Path("raw_data") / "2025 1023 01.xlsm"
SHEET = "Data to Graph"
CONC_COL = "B"          # Formula Fractions
SAMPLE_COL = "C"        # GH130 XL-40 R35 at 3 minutes
FIRST_ROW = 5
LAST_ROW = 52           # 48 formula-fraction rows (B5:B52)


def read_column(ws, col, r0, r1):
    vals = []
    for r in range(r0, r1 + 1):
        v = ws[f"{col}{r}"].value
        vals.append(v)
    return vals


def main():
    wb = load_workbook(WB, data_only=True)
    ws = wb[SHEET]

    label = ws[f"{SAMPLE_COL}4"].value or "Sample"
    conc = read_column(ws, CONC_COL, FIRST_ROW, LAST_ROW)
    pct_t = read_column(ws, SAMPLE_COL, FIRST_ROW, LAST_ROW)
    wb.close()

    # keep only rows where both concentration and %T are present
    pairs = [(c, t) for c, t in zip(conc, pct_t)
             if isinstance(c, (int, float)) and isinstance(t, (int, float))]
    conc = np.array([p[0] for p in pairs])
    pct_t = np.array([p[1] for p in pairs])

    print(f"Sample: {label}")
    print(f"Points: {len(conc)}  conc range {conc.min():.4f}-{conc.max():.4f}")

    # AUC over the mid dilution window as an example
    window = (0.10, 0.50)
    out = Path("example_absorbance_auc.png")
    _, _, auc = plot_absorbance_vs_concentration(
        conc, pct_t, sample_label=str(label), auc_window=window,
        kind="spline", out_path=str(out),
    )
    print(f"AUC over {window}: {auc['auc']:.4f}  (mean A {auc['mean_absorbance']:.3f})")
    print(f"Saved: {out.resolve()}")


if __name__ == "__main__":
    main()
