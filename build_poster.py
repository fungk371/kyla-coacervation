"""Wrap the verified poster PNG into a 16:9 PPTX slide (editable container)."""
import os
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])

png = os.path.join("Reports", "Coacervation_Poster.png")
with Image.open(png) as im:
    iw, ih = im.size
ar = iw / ih
sw, sh = prs.slide_width, prs.slide_height
if ar >= sw / sh:
    w = sw; h = int(sw / ar)
else:
    h = sh; w = int(sh * ar)
s.shapes.add_picture(png, (sw - w) // 2, (sh - h) // 2, width=w, height=h)

out = os.path.join("Reports", "Coacervation_Poster.pptx")
prs.save(out)
print("saved", out)
