"""Example time-course overlay from the real 'Data to Graph' GH130 XL-40 series."""
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from coacervation_plotting import plot_time_course

WB = Path("raw_data") / "2025 1023 01.xlsm"


def col(ws, letter, r0=5, r1=52):
    return [ws[f"{letter}{r}"].value for r in range(r0, r1 + 1)]


def clean(conc, y):
    pairs = [(c, v) for c, v in zip(conc, y)
             if isinstance(c, (int, float)) and isinstance(v, (int, float))]
    return np.array([p[0] for p in pairs]), np.array([p[1] for p in pairs])


def main():
    wb = load_workbook(WB, data_only=True)
    ws = wb["Data to Graph"]
    conc = col(ws, "B")
    # GH130 XL-40 R35 at 0/1.5/3/10.5/15 min -> columns E,F,G,H,I
    time_cols = {0: "E", 1.5: "F", 3: "G", 10.5: "H", 15: "I"}
    series = []
    for t, letter in time_cols.items():
        c, y = clean(conc, col(ws, letter))
        series.append({"time": t, "concentration": c, "pct_transmittance": y})
    wb.close()

    troughs = [{"label": "GH130 XL-40 R35", "series": series}]
    # AUC window 1:1 -> 1:10  => fractions 0.5 -> 1/11
    fig = plot_time_course(
        troughs, view="pct", auc_window=(1 / 11, 0.5),
        title="GH130 XL-40 R35 — 0–15 min Time Course",
        out_path="example_timecourse.png",
    )
    print("Saved example_timecourse.png")


if __name__ == "__main__":
    main()
