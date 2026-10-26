"""
Example: two-panel (%T and Absorbance) coacervation figure on fixed scales,
using two real GH130 samples from the 'Data to Graph' sheet.
"""

from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from coacervation_plotting import plot_transmittance_and_absorbance

WB = Path("raw_data") / "2025 1023 01.xlsm"


def col(ws, letter, r0=5, r1=52):
    return [ws[f"{letter}{r}"].value for r in range(r0, r1 + 1)]


def clean(conc, y):
    pairs = [(c, v) for c, v in zip(conc, y)
             if isinstance(c, (int, float)) and isinstance(v, (int, float))]
    return np.array([p[0] for p in pairs]), np.array([p[1] for p in pairs])


def main():
    wb = load_workbook(WB, data_only=True)
    ws = wb["Data to Graph"]
    conc = col(ws, "B")
    xl_label = ws["C4"].value      # GH130 XL-40 R35 at 3 minutes
    cg_label = ws["D4"].value      # GH130 CG17-40 R35 at 3 minutes
    xl = col(ws, "C")
    cg = col(ws, "D")
    wb.close()

    c1, xl = clean(conc, xl)
    c2, cg = clean(conc, cg)

    samples = [
        {"label": str(xl_label), "concentration": c1, "pct_transmittance": xl},
        {"label": str(cg_label), "concentration": c2, "pct_transmittance": cg},
    ]

    out = Path("example_dual_panel.png")
    _, _, results = plot_transmittance_and_absorbance(
        samples,
        auc_window=(0.10, 0.50),
        kind="spline",
        title="GH130 R35 — Coacervate Dissolution Profile (3 min)",
        pct_ylim=(0, 100),
        abs_ylim=(0, 0.6),
        out_path=str(out),
    )
    for r in results:
        print(f"{r['label']}: AUC={r['auc']:.4f}  meanA={r['mean_absorbance']:.3f}")
    print(f"Saved: {out.resolve()}")


if __name__ == "__main__":
    main()
